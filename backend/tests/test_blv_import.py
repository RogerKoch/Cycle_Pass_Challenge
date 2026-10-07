import openpyxl
import pytest

from backend.integrations.blv_import import import_blv
from backend.models.food import Food

HEADER = [
    "ID", "Name", "Synonyme", "Kategorie", "Energie, Kilokalorien (kcal)", "Herleitung des Wertes",
    "Fett, total (g)", "Kohlenhydrate, verfügbar (g)", "Protein (g)",
]


@pytest.fixture()
def blv_file(tmp_path):
    workbook = openpyxl.Workbook()
    generic = workbook.active
    generic.title = "Generische Lebensmittel"
    brands = workbook.create_sheet("Markenprodukte")
    for sheet in (generic, brands):
        sheet.append(["Schweizer Nährwertdatenbank"])
        sheet.append([])
        sheet.append(HEADER)
    generic.append([1, "Äpfel, roh", "Apfel", "Früchte", 54, "x", 0.1, 11.4, 0.3])
    generic.append([2, "Haselnuss", None, "Nüsse", 680, "x", 63.1, "Sp.", "<0.5"])
    brands.append([3, "Ovomaltine", None, "Getränke", 380, "x", 2.0, 80.0, 8.0])
    path = tmp_path / "blv.xlsx"
    workbook.save(path)
    return path


def test_imports_both_sheets_with_normalized_search_name(app, blv_file):
    assert import_blv(blv_file) == (3, 0)
    apple = Food.query.filter_by(source_id="1").one()
    assert apple.source == "blv"
    assert apple.kcal_100g == 54
    assert apple.search_name == "apfel, roh apfel"


def test_trace_and_below_limit_values_count_as_zero(app, blv_file):
    import_blv(blv_file)
    nut = Food.query.filter_by(source_id="2").one()
    assert (nut.carbs_100g, nut.protein_100g) == (0, 0)


def test_reimport_is_idempotent_and_keeps_own_portion(app, blv_file):
    import_blv(blv_file)
    apple = Food.query.filter_by(source_id="1").one()
    apple.portion_label, apple.portion_g = "Apfel", 150
    assert import_blv(blv_file) == (0, 3)
    assert Food.query.count() == 3
    assert Food.query.filter_by(source_id="1").one().portion_g == 150


def test_rejects_file_with_missing_column(app, tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Generische Lebensmittel"
    workbook.create_sheet("Markenprodukte")
    path = tmp_path / "broken.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        import_blv(path)


def test_text_numbers_are_read_and_duplicate_ids_update(app, tmp_path):
    workbook = openpyxl.Workbook()
    generic = workbook.active
    generic.title = "Generische Lebensmittel"
    brands = workbook.create_sheet("Markenprodukte")
    for sheet in (generic, brands):
        sheet.append(["Schweizer Nährwertdatenbank"])
        sheet.append([])
        sheet.append(HEADER)
    generic.append([1, "Quark", None, "Milch", "12,5", "x", 0.2, 3.0, 12.0])
    brands.append([1, "Quark", None, "Milch", 70, "x", 0.2, 3.0, 12.0])
    path = tmp_path / "dup.xlsx"
    workbook.save(path)
    assert import_blv(path) == (1, 1)
    assert Food.query.filter_by(source_id="1").one().kcal_100g == 70


def test_missing_sheet_raises_value_error(app, tmp_path):
    workbook = openpyxl.Workbook()
    generic = workbook.active
    generic.title = "Generische Lebensmittel"
    generic.append(["Schweizer Nährwertdatenbank"])
    generic.append([])
    generic.append(HEADER)
    path = tmp_path / "missing.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError, match="Markenprodukte"):
        import_blv(path)
